"""Independent CBAM calculator (not the app's engine): EC Guidance 3 / 4 / 5d formulas, values read
straight from the Commission's own files. Used as the expected-output oracle for the new test factory."""
from __future__ import annotations

from decimal import Decimal as D, getcontext
from pathlib import Path

import openpyxl

getcontext().prec = 40
RAW = str(Path(__file__).resolve().parents[3] / "backend" / "refdata" / "raw") + "/"  # the EC files, as downloaded

EF_NATURAL_GAS = D("56.1")       # t CO2 / TJ (NCV), IPCC 2006 / CBAM Implementing Reg. Annex VIII
MMBTU_TO_GJ = D("1.055056")      # GJ per MMBtu
NCV_PER_GCV = D("0.9")           # natural gas, IPCC 2006 Vol. 2 ch. 1
KCAL_TO_KJ = D("4.1868")
MARKUP_2026 = D("0.10")          # Guidance 3 p. 114-116
CBAM_FACTOR_2026 = D("0.975")    # Guidance 4 Table 2-1
CSCF_2026 = D("1")               # Implementing Decision (EU) 2026/1862


def india_default(cn: str) -> D:
    """Direct default value for India, read from the EC file (most specific CN prefix row)."""
    ws = openpyxl.load_workbook(RAW + "default_values_corrected_20260806.xlsx", read_only=True, data_only=True)["India"]
    best = None
    for row in ws.iter_rows(values_only=True):
        code = str(row[0] or "").replace(" ", "")
        if code and cn.startswith(code) and row[2] not in (None, "", "see below", "-", "N/A"):
            if best is None or len(code) > len(best[0]):
                best = (code, row[2], row[5])
    assert best, cn
    return D(str(best[1]).replace(",", "."))


def benchmarks(cn8: str) -> dict:
    """Column A and the column B values (by route/indicator) for an 8-digit CN, from the EC file."""
    ws = openpyxl.load_workbook(RAW + "cbam_benchmarks_20260206.xlsx", read_only=True, data_only=True)["Benchmarks"]
    out, on = {"B": {}}, False
    for row in ws.iter_rows(values_only=True):
        if row[0]:
            on = str(row[0]) == cn8
            if on:
                out["A"] = D(str(row[2]))
                out["B"][str(row[5] or "")] = D(str(row[4]))
        elif on and row[4] is not None:
            out["B"][str(row[5] or "")] = D(str(row[4]))
    assert "A" in out, cn8
    return out


def gj_ncv_from_mmbtu(mmbtu: D) -> D:
    return mmbtu * MMBTU_TO_GJ * NCV_PER_GCV


def see_complex(direct_t: D, al_t: D, precursors: list[tuple[D, D]]) -> dict:
    """SEE of a complex good = (AttrEm + sum M_i * SEE_i) / AL  (Guidance 3 section 6.2; 5d Example 3)."""
    direct = direct_t / al_t
    prec = sum((m * see for m, see in precursors), D(0)) / al_t
    return {"direct": direct, "precursor": prec, "see": direct + prec}


if __name__ == "__main__":
    # --- self-test against the Commission's worked examples (Guidance 5d, 14 Aug 2026) ---
    # Example 3, carbon-steel screws and nuts: AL 17,000 t, rod 20,000 t, rod SEE 1.567, gas 3.5 GJ(NCV)/t
    e = see_complex(D("3.5") * 17000 * EF_NATURAL_GAS / 1000, D(17000), [(D(20000), D("1.567"))])
    print("Ex3 carbon : direct %.4f precursor %.4f SEE %.4f m %.4f  (EC: 0.196 / 1.843 / 2.039 / 1.176)" % (e["direct"], e["precursor"], e["see"], D(20000) / 17000))
    e = see_complex(D("3.5") * 8200 * EF_NATURAL_GAS / 1000, D(8200), [(D(10000), D("1.783"))])
    print("Ex3 high-alloy: precursor %.4f SEE %.4f m %.4f  (EC: 2.175 / 2.371 / 1.220)" % (e["precursor"], e["see"], D(10000) / 8200))
    # Example 1, SEFA of rails = BM col A (rails) + m * SEFA(crude steel); 2027 factor 0.95
    sefa = D("0.057") + D("1.042") * D("1.156")
    print("Ex1 SEFA rails %.4f, 2027 %.4f  (EC: 1.261 / 1.198)" % (sefa, sefa * D("0.95")))
    print("official India DV 7213:", india_default("72139110"), " 7221:", india_default("72210010"), " 7318 16:", india_default("73181692"), india_default("73181639"))
    print("benchmarks 73181692:", benchmarks("73181692"), " 73181639:", benchmarks("73181639"))
    print("benchmarks 72139110:", benchmarks("72139110"), " 72210010:", benchmarks("72210010"))
