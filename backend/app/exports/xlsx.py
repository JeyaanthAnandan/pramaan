"""M10: Excel workbook - inputs, factors, calculations, evidence index."""

from __future__ import annotations

import io
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.evidence.models import Document, Evidence, Project

HEAD = PatternFill("solid", fgColor="EDEAE0")
STATUS_FILL = {
    "VERIFIED_2R": "DDF1E4",
    "APPROVED": "FDF0D5",
    "NEEDS_REVIEW": "FDE7C7",
    "MISSING": "F9D6D2",
    "DEFAULT_USED": "F9D6D2",
    "ASSUMPTION": "E8E6F6",
    "REFDATA": "E7EEF7",
}


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list[Any]], widths: list[int] | None = None) -> Any:
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEAD
    for r in rows:
        ws.append([float(x) if isinstance(x, Decimal) else x for x in r])
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def _num(v: str | None) -> float | str | None:
    if v in (None, ""):
        return None
    try:
        return float(Decimal(str(v)))
    except Exception:
        return v


def render(project: Project, result: dict[str, Any], evidence: list[Evidence], documents: list[Document]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Pramaan CBAM workbook — PROVISIONAL, not verified"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Operator", project.company])
    ws.append(["Period", result["period_label"]])
    ws.append(["Computed", result["computed_at"], "Run", result.get("run_id"), "Mode", result.get("mode")])
    ws.append(["Readiness %", _num(result["readiness"]["score"])])
    ws.append([])
    ws.append(
        [
            "CN",
            "Good",
            "AL t",
            "SEE t CO2/t",
            "Direct",
            "m",
            "Precursor data",
            "SEFA",
            "SEFA 2026",
            "Exported t",
            "Certificates 2026",
            "Buyer cost EUR (your data)",
            "Buyer cost EUR (EU default)",
            "Cheaper for the buyer",
        ]
    )
    for c in ws[7]:
        c.font = Font(bold=True)
        c.fill = HEAD
    for p in result["processes"]:
        if "see" in p:
            ws.append(
                [
                    p["cn_display"],
                    p["name"],
                    _num(p["al_t"]),
                    _num(p["see"]),
                    _num(p["direct"]),
                    _num(p["m"]),
                    p["precursor"]["data"],
                    _num(p["sefa"]),
                    _num(p["sefa_y"]),
                    _num(p["exports_t"]),
                    _num(p["certificates_2026"]),
                    _num(p["cost_2026_eur"]),
                    _num(p["buyer"]["cost_default_eur"]),
                    "your data" if p["buyer"]["basis_2026"] == "actual" else "EU default",
                ]
            )
    for i, w in enumerate([14, 30, 10, 12, 10, 10, 18, 10, 10, 11, 16, 18, 18, 18], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    docs = {d.id: d for d in documents}
    _sheet(
        wb,
        "Inputs (evidence)",
        [
            "Evidence id",
            "Document",
            "Doc type",
            "Field",
            "Label",
            "Period",
            "Process",
            "Value as written",
            "Value",
            "Unit",
            "Status",
            "Reader 2",
            "Agree",
            "Extractor",
            "Second reader",
            "Approved by",
            "Approved at",
            "box_2d [ymin,xmin,ymax,xmax]",
        ],
        [
            [
                e.id,
                docs[e.doc_id].filename if e.doc_id in docs else e.doc_id,
                docs[e.doc_id].doc_type if e.doc_id in docs else "",
                e.field_key,
                e.label,
                e.period,
                e.process,
                e.value_text,
                _num(e.value_num),
                e.unit,
                e.status,
                e.second_reading.reading_value if e.second_reading else None,
                e.second_reading.agrees if e.second_reading else None,
                e.extractor_model,
                e.second_reader_model,
                e.approved_by,
                e.approved_at,
                str(e.box_2d),
            ]
            for e in sorted(evidence, key=lambda x: (x.doc_id, x.field_key, x.period or ""))
        ],
        [18, 30, 22, 18, 26, 10, 8, 18, 12, 10, 14, 14, 8, 20, 20, 12, 20, 24],
    )
    tree = result["calc_tree"]
    calc = _sheet(
        wb,
        "Calculations",
        ["Node", "Label", "Value", "Unit", "Formula", "Inputs", "Status", "Source"],
        [
            [
                k,
                n["label"],
                _num(n["value"]),
                n["unit"],
                n.get("formula"),
                ", ".join(n["children"]),
                n["status"],
                (n.get("source") or {}).get("locator") or (n.get("source") or {}).get("doc_title") or "",
            ]
            for k, n in tree.items()
        ],
        [26, 46, 14, 14, 36, 50, 14, 50],
    )
    for row in calc.iter_rows(min_row=2):
        fill = STATUS_FILL.get(str(row[6].value))
        if fill:
            row[6].fill = PatternFill("solid", fgColor=fill)
    _sheet(
        wb,
        "Factors (refdata)",
        ["Node", "Label", "Value", "Unit", "Status", "Locator", "Document", "Published"],
        [
            [
                k,
                n["label"],
                _num(n["value"]),
                n["unit"],
                (n.get("source") or {}).get("status"),
                (n.get("source") or {}).get("locator"),
                ((n.get("source") or {}).get("document") or {}).get("title"),
                ((n.get("source") or {}).get("document") or {}).get("published"),
            ]
            for k, n in tree.items()
            if (n.get("source") or {}).get("kind") in ("refdata", "assumption")
        ],
        [26, 46, 12, 16, 12, 60, 60, 12],
    )
    rows = []
    for pid, series in result["money"]["series"].items():
        for s in series:
            rows.append(
                [
                    pid,
                    s["year"],
                    _num(s["price_eur"]),
                    _num(s["see_actual"]),
                    _num(s["sefa_y"]),
                    _num(s["actual_eur_t"]),
                    _num(s["see_default"]),
                    _num(s["sefa_default_y"]),
                    _num(s["default_eur_t"]),
                    _num(s["gap_eur_t"]),
                    _num(s["gap_inr_t"]),
                ]
            )
    _sheet(
        wb,
        "Money on the table",
        [
            "Process",
            "Year",
            "Price EUR",
            "SEE actual",
            "SEFA_y",
            "Cost actual EUR/t",
            "SEE default",
            "SEFA default_y",
            "Cost default EUR/t",
            "Difference EUR/t",
            "Difference INR/t",
        ],
        rows,
        [9, 7, 10, 11, 10, 15, 11, 13, 15, 14, 14],
    )
    _sheet(
        wb,
        "Control checks",
        ["Id", "Check", "Formula", "Result", "Tolerance", "Status", "Reference"],
        [[c["id"], c["name"], c["formula"], c["result"], c["tolerance"], c["status"], c["reference"]] for c in result.get("checks", [])],
        [24, 44, 50, 30, 14, 8, 40],
    )
    _sheet(
        wb,
        "Evidence index",
        ["Document", "Type", "Title", "Languages", "Period", "SHA-256", "Uploaded"],
        [[d.filename, d.doc_type, d.title, ",".join(d.languages), " to ".join(d.period or []), d.sha256, d.uploaded_at] for d in documents],
        [34, 26, 40, 12, 24, 66, 22],
    )
    for wsx in wb.worksheets:
        for row in wsx.iter_rows(min_row=1, max_row=1):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
